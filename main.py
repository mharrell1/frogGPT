import os
import sys
import sqlite3
from flask import Flask, render_template, request, jsonify, session, g, send_file
from werkzeug.security import generate_password_hash, check_password_hash

# Initialize ADK frogGPT runner if study-agent is available
froggpt_runner = None
froggpt_session_service = None
froggpt_init_error = None

try:
    study_agent_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'study-agent')
    if study_agent_path not in sys.path:
        sys.path.insert(0, study_agent_path)
    
    # Load dotenv from study-agent with override=True
    from dotenv import load_dotenv
    load_dotenv(os.path.join(study_agent_path, '.env'), override=True)
    
    # Explicitly set mandatory environment variables for AI Studio API keys
    os.environ['GOOGLE_GENAI_USE_ENTERPRISE'] = 'False'
    os.environ['GOOGLE_GENAI_USE_VERTEXAI'] = 'False'

    # Backup the existing 'app' module to prevent conflict with study-agent/app/ package
    flask_app_module = sys.modules.get('app')

    # Load the study-agent/app package manually to resolve package name conflict
    import importlib.util
    study_agent_app_init = os.path.join(study_agent_path, 'app', '__init__.py')
    spec = importlib.util.spec_from_file_location("app", study_agent_app_init)
    study_agent_app_module = importlib.util.module_from_spec(spec)
    
    # Register the study-agent/app package as 'app' temporarily
    sys.modules['app'] = study_agent_app_module
    spec.loader.exec_module(study_agent_app_module)

    # Now we can import the agent and run Gunicorn without packages conflict!
    from app.agent import root_agent

    # Restore the original flask app module in sys.modules
    if flask_app_module:
        sys.modules['app'] = flask_app_module
    from google.adk.runners import Runner
    from google.adk.sessions import InMemorySessionService

    # Monkey patch to fix ADK 2.0 active task isolation scope bug
    import google.adk.runners as adk_runners
    def patched_find_active_task_isolation_scope(session):
        finished_scopes = set()
        for event in session.events:
            scope = event.isolation_scope
            if not scope:
                continue
            if event.content and event.content.parts:
                for part in event.content.parts:
                    fr = part.function_response
                    if fr and fr.name == 'finish_task':
                        response = fr.response or {}
                        if response.get('result') == 'Task completed.':
                            finished_scopes.add(scope)
                        break
        for event in reversed(session.events):
            scope = event.isolation_scope
            if scope and scope not in finished_scopes:
                return scope
        return None
    adk_runners._find_active_task_isolation_scope = patched_find_active_task_isolation_scope

    froggpt_session_service = InMemorySessionService()
    froggpt_runner = Runner(agent=root_agent, session_service=froggpt_session_service, app_name="frogGPT")
    print("Successfully initialized frogGPT ADK Runner.")
except Exception as e:
    import traceback
    adk_ver = "unknown"
    try:
        import importlib.metadata
        adk_ver = importlib.metadata.version('google-adk')
    except Exception:
        pass
    froggpt_init_error = f"[adk-version: {adk_ver}] " + traceback.format_exc()
    traceback.print_exc()
# Initialize Google GenAI client directly using API key
genai_client = None
try:
    from google import genai
    from google.genai import types
    api_key = os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")
    if api_key:
        genai_client = genai.Client(api_key=api_key)
    else:
        genai_client = genai.Client()
    print("Successfully initialized Google GenAI Client.")
except Exception as e:
    print(f"Failed to initialize Google GenAI Client: {e}")

# ---------------------------------------------------------------------------
# Monkey patch google-genai Client models service to track exact API call count
# ---------------------------------------------------------------------------
try:
    from google.genai.models import Models, AsyncModels
    from flask import has_request_context, g

    def track_api_call():
        if has_request_context():
            g.gemini_calls_count = getattr(g, 'gemini_calls_count', 0) + 1

    # Patched generate_content
    orig_generate_content = Models.generate_content
    def patched_generate_content(self, *args, **kwargs):
        track_api_call()
        return orig_generate_content(self, *args, **kwargs)
    Models.generate_content = patched_generate_content

    # Patched generate_content_stream
    orig_generate_content_stream = Models.generate_content_stream
    def patched_generate_content_stream(self, *args, **kwargs):
        track_api_call()
        return orig_generate_content_stream(self, *args, **kwargs)
    Models.generate_content_stream = patched_generate_content_stream

    # Patched async generate_content
    orig_async_generate_content = AsyncModels.generate_content
    async def patched_async_generate_content(self, *args, **kwargs):
        track_api_call()
        return await orig_async_generate_content(self, *args, **kwargs)
    AsyncModels.generate_content = patched_async_generate_content

    # Patched async generate_content_stream
    orig_async_generate_content_stream = AsyncModels.generate_content_stream
    async def patched_async_generate_content_stream(self, *args, **kwargs):
        track_api_call()
        return await orig_async_generate_content_stream(self, *args, **kwargs)
    AsyncModels.generate_content_stream = patched_async_generate_content_stream

    print("Successfully monkey-patched google-genai Models to track quota.")
except Exception as e:
    print(f"Failed to monkey-patch google-genai Models: {e}")

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'froggpt-secret-key-12345')
app.config['SESSION_COOKIE_NAME'] = 'froggpt_session'
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'

DATABASE = os.environ.get('DB_PATH', 'tasks.db')

def get_db():
    db = getattr(g, '_database', None)
    if db is None:
        db_dir = os.path.dirname(DATABASE)
        if db_dir and not os.path.exists(db_dir):
            os.makedirs(db_dir, exist_ok=True)
        db = g._database = sqlite3.connect(DATABASE)
        db.row_factory = sqlite3.Row
    return db

@app.teardown_appcontext
def close_connection(exception):
    db = getattr(g, '_database', None)
    if db is not None:
        db.close()

def init_db():
    db = get_db()
    # Create users table
    db.execute('''
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL
        )
    ''')
    # Create tasks table if not exists
    db.execute('''
        CREATE TABLE IF NOT EXISTS tasks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            completed INTEGER DEFAULT 0,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    # Create pomodoro_logs table
    db.execute('''
        CREATE TABLE IF NOT EXISTS pomodoro_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER REFERENCES users(id),
            category TEXT NOT NULL,
            duration_minutes INTEGER NOT NULL,
            completed_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    # Add columns to tasks table if they don't exist
    cursor = db.execute("PRAGMA table_info(tasks)")
    columns = [row['name'] for row in cursor.fetchall()]
    if 'user_id' not in columns:
        db.execute("ALTER TABLE tasks ADD COLUMN user_id INTEGER REFERENCES users(id)")
    if 'due_date' not in columns:
        db.execute("ALTER TABLE tasks ADD COLUMN due_date TEXT")
    if 'category' not in columns:
        db.execute("ALTER TABLE tasks ADD COLUMN category TEXT")
    if 'urgency' not in columns:
        db.execute("ALTER TABLE tasks ADD COLUMN urgency TEXT")
    if 'notes' not in columns:
        db.execute("ALTER TABLE tasks ADD COLUMN notes TEXT")
    if 'completed_at' not in columns:
        db.execute("ALTER TABLE tasks ADD COLUMN completed_at TIMESTAMP")
    db.commit()

with app.app_context():
    init_db()

@app.route('/')
def index():
    return render_template('index.html')

# Authentication APIs
@app.route('/api/register', methods=['POST'])
def register():
    data = request.get_json() or {}
    username = data.get('username', '').strip()
    password = data.get('password', '')

    if not username or not password:
        return jsonify({'error': 'Username and password are required'}), 400

    if len(password) < 6:
        return jsonify({'error': 'Password must be at least 6 characters'}), 400

    db = get_db()
    try:
        # Check if user already exists
        user = db.execute('SELECT * FROM users WHERE username = ?', (username,)).fetchone()
        if user:
            return jsonify({'error': 'Username already taken'}), 400

        password_hash = generate_password_hash(password)
        cursor = db.execute('INSERT INTO users (username, password) VALUES (?, ?)', (username, password_hash))
        db.commit()
        user_id = cursor.lastrowid

        session['user_id'] = user_id
        session['username'] = username
        return jsonify({'success': True, 'username': username})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/login', methods=['POST'])
def login():
    data = request.get_json() or {}
    username = data.get('username', '').strip()
    password = data.get('password', '')

    if not username or not password:
        return jsonify({'error': 'Username and password are required'}), 400

    db = get_db()
    user = db.execute('SELECT * FROM users WHERE username = ?', (username,)).fetchone()
    if not user or not check_password_hash(user['password'], password):
        return jsonify({'error': 'Invalid username or password'}), 401

    session['user_id'] = user['id']
    session['username'] = user['username']
    return jsonify({'success': True, 'username': user['username']})

@app.route('/api/logout', methods=['POST'])
def logout():
    session.clear()
    return jsonify({'success': True})

@app.route('/api/user', methods=['GET'])
def get_user():
    if 'user_id' in session:
        return jsonify({'logged_in': True, 'username': session['username']})
    return jsonify({'logged_in': False})

# Task APIs
@app.route('/api/tasks', methods=['GET'])
def get_tasks():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    db = get_db()
    tasks = db.execute('SELECT * FROM tasks WHERE user_id = ? ORDER BY id DESC', (session['user_id'],)).fetchall()
    
    return jsonify([
        {
            'id': task['id'],
            'title': task['title'],
            'completed': bool(task['completed']),
            'created_at': task['created_at'],
            'due_date': task['due_date'],
            'category': task['category'] or 'School',
            'urgency': task['urgency'] or 'medium',
            'notes': task['notes'] or ''
        } for task in tasks
    ])

@app.route('/api/tasks', methods=['POST'])
def add_task():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401

    data = request.get_json() or {}
    title = (data.get('title') or '').strip()
    due_date = (data.get('due_date') or '').strip() or None
    category = (data.get('category') or '').strip() or 'School'
    urgency = (data.get('urgency') or '').strip() or 'medium'
    notes = (data.get('notes') or '').strip() or None

    if not title:
        return jsonify({'error': 'Task title is required'}), 400

    db = get_db()
    try:
        cursor = db.execute(
            'INSERT INTO tasks (user_id, title, completed, due_date, category, urgency, notes) VALUES (?, ?, 0, ?, ?, ?, ?)',
            (session['user_id'], title, due_date, category, urgency, notes)
        )
        db.commit()
        task_id = cursor.lastrowid
        
        task = db.execute('SELECT * FROM tasks WHERE id = ?', (task_id,)).fetchone()
        return jsonify({
            'id': task['id'],
            'title': task['title'],
            'completed': bool(task['completed']),
            'created_at': task['created_at'],
            'due_date': task['due_date'],
            'category': task['category'] or 'School',
            'urgency': task['urgency'] or 'medium',
            'notes': task['notes'] or ''
        }), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/tasks/<int:task_id>', methods=['PUT'])
def update_task(task_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401

    data = request.get_json() or {}
    completed = data.get('completed')
    title = data.get('title')
    due_date = data.get('due_date')
    category = data.get('category')
    urgency = data.get('urgency')
    notes = data.get('notes')

    db = get_db()
    task = db.execute('SELECT * FROM tasks WHERE id = ? AND user_id = ?', (task_id, session['user_id'])).fetchone()
    if not task:
        return jsonify({'error': 'Task not found'}), 404

    try:
        update_fields = []
        params = []
        if completed is not None:
            import datetime
            update_fields.append("completed = ?")
            params.append(1 if completed else 0)
            update_fields.append("completed_at = ?")
            params.append(datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S') if completed else None)
        if title is not None:
            update_fields.append("title = ?")
            params.append((title or '').strip())
        if due_date is not None:
            update_fields.append("due_date = ?")
            params.append((due_date or '').strip() or None)
        if category is not None:
            update_fields.append("category = ?")
            params.append((category or '').strip() or 'School')
        if urgency is not None:
            update_fields.append("urgency = ?")
            params.append((urgency or '').strip() or 'medium')
        if notes is not None:
            update_fields.append("notes = ?")
            params.append((notes or '').strip() or None)

        if not update_fields:
            return jsonify({'error': 'No fields to update'}), 400

        params.append(task_id)
        query = f"UPDATE tasks SET {', '.join(update_fields)} WHERE id = ?"
        db.execute(query, tuple(params))
        db.commit()
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/tasks/<int:task_id>', methods=['DELETE'])
def delete_task(task_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401

    db = get_db()
    task = db.execute('SELECT * FROM tasks WHERE id = ? AND user_id = ?', (task_id, session['user_id'])).fetchone()
    if not task:
        return jsonify({'error': 'Task not found'}), 404

    try:
        db.execute('DELETE FROM tasks WHERE id = ?', (task_id,))
        db.commit()
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

def run_async(coro):
    import asyncio
    try:
        loop = asyncio.get_event_loop()
    except RuntimeError:
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
    return loop.run_until_complete(coro)



@app.route('/api/froggpt/import', methods=['POST'])
def froggpt_import_file():
    if 'file' not in request.files:
        return jsonify({'error': 'No file part in request'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
        
    filename = file.filename
    ext = os.path.splitext(filename)[1].lower()
    
    if ext not in ['.txt', '.md', '.docx', '.pdf']:
        return jsonify({'error': 'Unsupported file format. Please upload .txt, .md, .docx, or .pdf files.'}), 400
        
    try:
        content = ""
        if ext in ['.txt', '.md']:
            content = file.read().decode('utf-8', errors='ignore')
        elif ext == '.docx':
            import docx
            from io import BytesIO
            doc = docx.Document(BytesIO(file.read()))
            paragraphs = [p.text for p in doc.paragraphs]
            content = "\n".join(paragraphs)
        elif ext == '.pdf':
            import pypdf
            from io import BytesIO
            reader = pypdf.PdfReader(BytesIO(file.read()))
            text_list = []
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text_list.append(t)
            content = "\n".join(text_list)
            
        return jsonify({
            'success': True,
            'filename': filename,
            'content': content,
            'char_count': len(content)
        })
    except Exception as e:
        return jsonify({'error': f'Failed to parse file: {str(e)}'}), 500

# frogGPT AI Study Agent API
@app.route('/api/froggpt', methods=['POST'])
def froggpt_chat():
    if not froggpt_runner or not froggpt_session_service:
        err_msg = 'frogGPT is not initialized properly. Check server logs and environment variables.'
        if froggpt_init_error:
            err_msg += f" Details: {froggpt_init_error}"
        return jsonify({'error': err_msg}), 500

    data = request.get_json() or {}
    message_text = data.get('message', '').strip()
    if not message_text:
        return jsonify({'error': 'Message is required'}), 400

    orig_api_clients = {}

    # Get or create session ID for the user
    user_id = session.get('username', 'guest_user')
    session_id = data.get('session_id')
    
    if session_id:
        try:
            sess = froggpt_session_service.get_session_sync(user_id=user_id, session_id=session_id, app_name="frogGPT")
            if not sess:
                session_id = None
        except Exception:
            session_id = None

    if not session_id:
        try:
            new_sess = froggpt_session_service.create_session_sync(user_id=user_id, app_name="frogGPT")
            session_id = new_sess.id
        except Exception as e:
            return jsonify({'error': f'Failed to create session: {e}'}), 500



    try:
        from google.genai import types
        from google.adk.agents.run_config import RunConfig, StreamingMode
        
        model_name = data.get('model', 'gemini-2.5-flash').strip()
        imported_content = data.get('imported_content', '').strip()
        


        # 2. Gemini Pipeline (Default ADK)
        # Prepend imported notes if present
        if imported_content:
            message_text = f"Notes:\n{imported_content}\n\nRequest: {message_text}"

        allowed_models = ['gemini-2.5-flash', 'gemini-2.5-pro', 'gemini-1.5-flash', 'gemini-1.5-pro']
        if model_name not in allowed_models:
            model_name = 'gemini-2.5-flash'
            
        agents_to_update = [root_agent]
        if hasattr(root_agent, 'sub_agents') and root_agent.sub_agents:
            agents_to_update.extend(root_agent.sub_agents)

        for ag in agents_to_update:
            if hasattr(ag, 'model') and ag.model:
                if hasattr(ag.model, 'model'):
                    ag.model.model = model_name

        sess = froggpt_session_service.get_session_sync(user_id=user_id, session_id=session_id, app_name="frogGPT")
        if sess:
            sess.state["user_id"] = user_id

        msg = types.Content(role="user", parts=[types.Part.from_text(text=message_text)])
        events = list(
            froggpt_runner.run(
                new_message=msg,
                user_id=user_id,
                session_id=session_id,
                run_config=RunConfig(streaming_mode=StreamingMode.SSE),
            )
        )
        
        # Check if any sub-agent successfully produced a structured output dict
        subagent_markdown = ""
        structured_data = None
        subagent_author = None
        try:
            from app.tools import (
                format_flashcards_markdown,
                format_study_guide_markdown,
                format_quiz_markdown,
                format_practice_test_markdown,
                format_explanation_markdown,
            )
            import json
            
            for event in events:
                output_val = getattr(event, 'output', None)
                if output_val is not None:
                    # Convert Pydantic model to dict if necessary
                    output_dict = None
                    if isinstance(output_val, dict):
                        output_dict = output_val
                    elif hasattr(output_val, 'model_dump'):
                        output_dict = output_val.model_dump()
                    elif hasattr(output_val, 'dict'):
                        output_dict = output_val.dict()
                    
                    if output_dict is not None and isinstance(output_dict, dict):
                        author = getattr(event, 'author', '')
                        out_json = json.dumps(output_dict)
                        res = {}
                        
                        # 1. Prioritize strict author matching
                        if author == 'quiz_agent':
                            res = format_quiz_markdown(out_json)
                            if res.get('status') == 'success':
                                structured_data = output_dict
                                subagent_author = 'quiz_agent'
                        elif author == 'flashcard_agent':
                            res = format_flashcards_markdown(out_json)
                            if res.get('status') == 'success':
                                structured_data = output_dict
                                subagent_author = 'flashcard_agent'
                        elif author == 'study_guide_agent':
                            res = format_study_guide_markdown(out_json)
                            if res.get('status') == 'success':
                                structured_data = output_dict
                                subagent_author = 'study_guide_agent'
                        elif author == 'test_agent':
                            res = format_practice_test_markdown(out_json)
                            if res.get('status') == 'success':
                                structured_data = output_dict
                                subagent_author = 'test_agent'
                        elif author == 'explain_agent':
                            res = format_explanation_markdown(out_json)
                            if res.get('status') == 'success':
                                structured_data = output_dict
                                subagent_author = 'explain_agent'
                                
                        # 2. Fallbacks (only if author was not matched above)
                        if not subagent_author or res.get('status') != 'success':
                            if 'questions' in output_dict and any('options' in str(q) for q in output_dict.get('questions', [])):
                                res = format_quiz_markdown(out_json)
                                if res.get('status') == 'success':
                                    structured_data = output_dict
                                    subagent_author = 'quiz_agent'
                            elif 'cards' in output_dict:
                                res = format_flashcards_markdown(out_json)
                                if res.get('status') == 'success':
                                    structured_data = output_dict
                                    subagent_author = 'flashcard_agent'
                            elif 'sections' in output_dict:
                                res = format_study_guide_markdown(out_json)
                                if res.get('status') == 'success':
                                    structured_data = output_dict
                                    subagent_author = 'study_guide_agent'
                            elif 'true_false' in output_dict or 'multiple_choice' in output_dict:
                                res = format_practice_test_markdown(out_json)
                                if res.get('status') == 'success':
                                    structured_data = output_dict
                                    subagent_author = 'test_agent'
                            elif 'simple_explanation' in output_dict or 'analogy' in output_dict:
                                res = format_explanation_markdown(out_json)
                                if res.get('status') == 'success':
                                    structured_data = output_dict
                                    subagent_author = 'explain_agent'
                        
                        if res.get('status') == 'success' and res.get('markdown'):
                            subagent_markdown = res['markdown']
                            break
        except Exception as fmt_err:
            print(f"Debug - error formatting subagent output: {fmt_err}")

        # First try to find the final aggregated event
        full_text = subagent_markdown
        if not full_text:
            for event in events:
                if not getattr(event, 'partial', False) and getattr(event, 'content', None) and getattr(event.content, 'parts', None):
                    text_parts = [part.text for part in event.content.parts if getattr(part, 'text', None)]
                    if text_parts:
                        full_text = "".join(text_parts)
                    
        # If no final aggregated event found, concatenate all partial events
        if not full_text:
            for event in events:
                if getattr(event, 'partial', True) and getattr(event, 'content', None) and getattr(event.content, 'parts', None):
                    for part in event.content.parts:
                        if getattr(part, 'text', None):
                            full_text += part.text
                        
        if not full_text:
            print(f"Debug - events received: {events}")
            # Try to extract any error or quota information from the events
            err_msg = ""
            for event in events:
                if getattr(event, 'error', None):
                    err_msg += str(event.error) + " "
                elif getattr(event, 'error_message', None):
                    err_msg += str(event.error_message) + " "
                elif getattr(event, 'status', None) and str(event.status) != "OK":
                    err_msg += str(event.status) + " "
            
            events_str = str(events)
            if "not found" in events_str.lower() or "not supported" in events_str.lower():
                full_text = f"⚠️ **Model Not Supported or Not Found**\n\n🐸 Lily tried to use the selected model, but it is not available on your API key or is not supported.\n\n**Details**: `{err_msg or 'Model not supported'}`\n\n🔄 **Fix**: Please switch your study model to **Gemini 2.5 Flash** or **Gemini 2.5 Pro** in the dropdown select!"
            elif err_msg or "429" in events_str or "quota" in events_str.lower() or "resource_exhausted" in events_str.lower():
                full_text = f"⚠️ **Google AI Studio Quota Exceeded / Rate Limit**\n\n🐸 Lily tried to answer, but the Gemini API free tier quota was reached (`429 Too Many Requests`).\n\n**Details**: {err_msg or 'Free tier limits requests to 20 per minute or daily limits.'}\n\n⏳ **Fix**: Please wait about 45–60 seconds, take a deep breath, and try clicking the button again!"
            else:
                full_text = f"⚠️ **Google AI Studio Rate Limit / No Response**\n\n🐸 Lily couldn't generate a text response. This usually happens when the Gemini API free tier rate limit (`429 Resource Exhausted`) is reached during a multi-agent workflow.\n\n`Debug Info: {events_str}`\n\n⏳ **Fix**: Please wait about 60 seconds and try asking again!"

        final_response_text = getattr(g, 'fallback_prefix', '') + full_text

        return jsonify({
            'success': True,
            'response': final_response_text,
            'session_id': session_id,
            'structured_data': structured_data,
            'subagent_author': subagent_author,
            'calls_made': getattr(g, 'gemini_calls_count', 0)
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        # Check if the exception itself is a quota error
        err_str = str(e)
        if "429" in err_str or "quota" in err_str.lower() or "resource_exhausted" in err_str.lower():
            return jsonify({'success': True, 'response': f"⚠️ **Google AI Studio Quota Exceeded (`429`)**\n\n🐸 Lily tried to answer, but the Gemini API free tier quota was reached.\n\n**Error Details**: {err_str}\n\n⏳ **Fix**: Please wait about 45–60 seconds, take a deep breath, and try again!", 'session_id': session_id, 'calls_made': getattr(g, 'gemini_calls_count', 0)})
        return jsonify({'error': str(e)}), 500
    finally:
        # Restore original cached clients
        for model, orig_client in orig_api_clients.items():
            if orig_client is None:
                model.__dict__.pop('api_client', None)
            else:
                model.__dict__['api_client'] = orig_client

@app.route('/api/froggpt/history', methods=['GET'])
def froggpt_list_history():
    user_id = session.get('username')
    if not user_id:
        return jsonify({'error': 'You must be logged in to access history.'}), 401

    if not froggpt_session_service:
        return jsonify({'error': 'Session service not initialized'}), 500

    try:
        resp = froggpt_session_service.list_sessions_sync(app_name="frogGPT", user_id=user_id)
        sessions_list = []
        if resp and resp.sessions:
            for s in resp.sessions:
                # Format update time
                import datetime
                dt_str = "Unknown date"
                if getattr(s, 'last_update_time', None):
                    try:
                        dt_str = datetime.datetime.fromtimestamp(s.last_update_time).strftime("%b %d, %Y %I:%M %p")
                    except Exception:
                        pass
                
                # Fetch full session once to extract the first message as summary
                first_msg = "New Study Session"
                try:
                    full_sess = froggpt_session_service.get_session_sync(user_id=user_id, session_id=s.id, app_name="frogGPT")
                    if full_sess and full_sess.events:
                        for ev in full_sess.events:
                            if ev.author == 'user' and ev.content and ev.content.parts:
                                txt = "".join([p.text for p in ev.content.parts if getattr(p, 'text', None)]).strip()
                                if txt:
                                    first_msg = txt[:45] + ("..." if len(txt) > 45 else "")
                                    break
                except Exception:
                    pass

                sessions_list.append({
                    'id': s.id,
                    'updated_at': dt_str,
                    'summary': first_msg
                })
        # Sort in reverse chronological order
        sessions_list.reverse()
        return jsonify({'success': True, 'sessions': sessions_list})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/froggpt/history/<session_id>', methods=['GET'])
def froggpt_session_history(session_id):
    user_id = session.get('username')
    if not user_id:
        return jsonify({'error': 'You must be logged in to access history.'}), 401
    
    if not froggpt_session_service:
        return jsonify({'error': 'Session service not initialized'}), 500
        
    try:
        sess = froggpt_session_service.get_session_sync(user_id=user_id, session_id=session_id, app_name="frogGPT")
        if not sess:
            return jsonify({'error': 'Session not found'}), 404
            
        history = []
        if sess.events:
            for ev in sess.events:
                text = ""
                if ev.content and ev.content.parts:
                    text = "".join([p.text for p in ev.content.parts if getattr(p, 'text', None)])
                
                # Check for structured JSON blocks
                import re
                import json
                structured_data = None
                subagent_author = None
                
                json_match = re.search(r'```json\s*(.*?)\s*```', text, re.DOTALL)
                if json_match:
                    try:
                        structured_data = json.loads(json_match.group(1).strip())
                        if 'cards' in structured_data:
                            subagent_author = 'flashcard_agent'
                        elif 'difficulty' in structured_data or ('questions' in structured_data and 'options' in str(structured_data)):
                            subagent_author = 'quiz_agent'
                        elif 'true_false' in structured_data or 'multiple_choice' in structured_data:
                            subagent_author = 'test_agent'
                        elif 'sections' in structured_data:
                            subagent_author = 'study_guide_agent'
                            
                        # Keep the text clean without the raw json match
                        text = text.replace(json_match.group(0), "").strip()
                    except Exception:
                        pass

                history.append({
                    'author': ev.author,
                    'text': text,
                    'structured_data': structured_data,
                    'subagent_author': subagent_author,
                    'timestamp': getattr(ev, 'timestamp', None)
                })
        return jsonify({'success': True, 'events': history})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# --- AI Note Taking Agent endpoints ---
@app.route('/api/notes/transcribe', methods=['POST'])
def notes_transcribe():
    active_client = genai_client
    if not active_client:
        return jsonify({'error': 'Gemini API client not initialized. Ensure GEMINI_API_KEY is configured.'}), 500

    if 'audio' not in request.files:
        return jsonify({'error': 'No audio file provided'}), 400

    audio_file = request.files['audio']
    if audio_file.filename == '':
        return jsonify({'error': 'No selected audio file'}), 400

    try:
        audio_bytes = audio_file.read()
        mime_type = audio_file.mimetype or 'audio/webm'
        
        # Call Gemini to transcribe
        response = active_client.models.generate_content(
            model='gemini-2.5-flash',
            contents=[
                types.Part.from_bytes(
                    data=audio_bytes,
                    mime_type=mime_type
                ),
                "Please transcribe the following audio recording with extremely high fidelity. Output the transcription directly, word-for-word, without any additional explanations, notes, or introductions."
            ]
        )
        
        transcript = response.text or ""
        return jsonify({'success': True, 'transcript': transcript, 'calls_made': getattr(g, 'gemini_calls_count', 0)})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/notes/summarize', methods=['POST'])
def notes_summarize():
    active_client = genai_client
    if not active_client:
        return jsonify({'error': 'Gemini API client not initialized. Ensure GEMINI_API_KEY is configured.'}), 500

    data = request.get_json() or {}
    transcript = data.get('transcript', '').strip()
    note_type = data.get('type', 'lecture').strip().lower()

    if not transcript:
        return jsonify({'error': 'No transcript content provided'}), 400

    try:
        if note_type == 'meeting':
            prompt = (
                "Review the following transcript of a meeting. Generate highly structured meeting notes. "
                "Include:\n"
                "1. A concise overview/executive summary of the meeting.\n"
                "2. Key decisions made.\n"
                "3. Bullet points of main discussion items.\n"
                "4. A clear, actionable checklist of action items, specifying who is responsible if mentioned.\n\n"
                f"Transcript:\n{transcript}"
            )
        else:
            prompt = (
                "Review the following transcript of a lecture. Generate a structured lecture summary. "
                "Include:\n"
                "1. A general introduction of the core topic.\n"
                "2. Detailed key concepts and definitions explained.\n"
                "3. Crucial take-aways and insights in structured bullet points.\n"
                "4. Follow-up study questions or topics for research.\n\n"
                f"Transcript:\n{transcript}"
            )

        response = active_client.models.generate_content(
            model='gemini-2.5-flash',
            contents=prompt
        )
        notes = response.text or ""
        return jsonify({'success': True, 'notes': notes, 'calls_made': getattr(g, 'gemini_calls_count', 0)})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/notes/export', methods=['POST'])
def notes_export():
    data = request.get_json() or {}
    notes = data.get('notes', '').strip()
    filename = data.get('filename', 'my_notes').strip()
    export_format = data.get('format', 'pdf').strip().lower()

    if not notes:
        return jsonify({'error': 'No notes content provided'}), 400

    safe_filename = "".join(c if c.isalnum() or c in "-_ " else "_" for c in filename).strip()
    if not safe_filename:
        safe_filename = "notes_summary"

    try:
        # Import study-agent tools
        from app.tools import export_as_pdf, export_as_docx
        
        home = os.path.expanduser('~')
        local_notes_dir = os.path.join(home, 'Documents', 'Note-Agent', 'Notes')
        
        is_local = os.path.exists(os.path.join(home, 'Documents'))
        if is_local:
            target_dir = local_notes_dir
        else:
            target_dir = "/tmp"

        os.makedirs(target_dir, exist_ok=True)

        if export_format == 'pdf':
            res = export_as_pdf(notes, safe_filename, target_dir)
        elif export_format == 'docx':
            res = export_as_docx(notes, safe_filename, target_dir)
        else:
            return jsonify({'error': 'Invalid format. Supported formats: pdf, docx'}), 400

        if res.get('status') == 'success':
            file_path = res.get('file_path')
            
            session['last_export_path'] = file_path
            session['last_export_name'] = os.path.basename(file_path)
            
            return jsonify({
                'success': True,
                'file_path': file_path,
                'filename': os.path.basename(file_path),
                'saved_locally': is_local,
                'local_path': local_notes_dir if is_local else None
            })
        else:
            return jsonify({'error': res.get('file_path') or 'Export failed'}), 500

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/notes/download', methods=['GET'])
def notes_download():
    file_path = session.get('last_export_path')
    file_name = session.get('last_export_name', 'notes.pdf')
    
    if not file_path or not os.path.exists(file_path):
        return "No export found or file has expired. Please export the notes again.", 404
        
    return send_file(file_path, as_attachment=True, download_name=file_name)

@app.route('/api/notes/import-link', methods=['POST'])
def import_video_link():
    if not genai_client:
        return jsonify({'error': 'Gemini API client not initialized. Ensure GEMINI_API_KEY is configured.'}), 500

    data = request.get_json() or {}
    url = data.get('url', '').strip()

    if not url:
        return jsonify({'error': 'No URL provided'}), 400

    try:
        import tempfile
        import requests
        import time

        is_youtube = ('youtube.com' in url.lower() or 
                      'youtu.be' in url.lower() or 
                      (len(url) == 11 and re.match(r'^[a-zA-Z0-9_-]{11}$', url)))

        if is_youtube:
            # 1. Parse video ID
            import re
            video_id = None
            if len(url) == 11 and re.match(r'^[a-zA-Z0-9_-]{11}$', url):
                video_id = url
            else:
                patterns = [
                    r'v=([a-zA-Z0-9_-]{11})',
                    r'youtu\.be/([a-zA-Z0-9_-]{11})',
                    r'embed/([a-zA-Z0-9_-]{11})',
                    r'shorts/([a-zA-Z0-9_-]{11})',
                    r'live/([a-zA-Z0-9_-]{11})',
                    r'v/([a-zA-Z0-9_-]{11})',
                ]
                for pattern in patterns:
                    match = re.search(pattern, url)
                    if match:
                        video_id = match.group(1)
                        break
            
            if not video_id:
                return jsonify({'error': 'Could not parse YouTube Video ID from the link. Make sure it is a valid YouTube URL.'}), 400

            # 2. Try fetching transcript via youtube-transcript.ai API
            try:
                import requests
                import re
                yt_api_url = f"https://youtube-transcript.ai/transcript/{video_id}.txt"
                headers = {
                    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
                }
                response = requests.get(yt_api_url, headers=headers, timeout=15)
                if response.status_code == 200 and "## Transcript" in response.text:
                    lines = response.text.split('\n')
                    transcript_lines = []
                    in_transcript = False
                    for line in lines:
                        if line.startswith('## Transcript'):
                            in_transcript = True
                            continue
                        if in_transcript:
                            line_clean = line.strip()
                            if line_clean:
                                # Remove timestamps like [0:01] or [12:34:56]
                                line_clean = re.sub(r'^\[\d+(?::\d+)+\]\s*', '', line_clean)
                                transcript_lines.append(line_clean)
                    
                    if transcript_lines:
                        transcript = "\n".join(transcript_lines)
                        return jsonify({'success': True, 'transcript': transcript, 'calls_made': getattr(g, 'gemini_calls_count', 0)})
                
                raise Exception("API did not return a valid transcript structure (e.g. rate-limited).")
            except Exception as yt_err:
                print(f"youtube-transcript.ai API failed for {video_id}: {yt_err}")
                # Fallback: Process the YouTube URL directly using Gemini's native multimodal capabilities
                try:
                    normalized_url = f"https://www.youtube.com/watch?v={video_id}"
                    youtube_video = types.Part.from_uri(
                        file_uri=normalized_url,
                        mime_type='video/mp4'
                    )
                    response = genai_client.models.generate_content(
                        model='gemini-2.5-flash',
                        contents=[
                            youtube_video,
                            "Transcribe the audio of this video as accurately and completely as possible. Only return the transcript."
                        ]
                    )
                    transcript = response.text or ""
                    if not transcript:
                        raise Exception("Gemini returned an empty transcript.")
                    return jsonify({'success': True, 'transcript': transcript, 'calls_made': getattr(g, 'gemini_calls_count', 0)})
                except Exception as gemini_err:
                    print(f"Direct YouTube URL transcription failed: {gemini_err}")
                    # Final Fallback: Download audio stream using yt-dlp (in case they have proxy config or run locally)
                    try:
                        import yt_dlp
                        import tempfile
                        import glob
                        
                        temp_dir = tempfile.gettempdir()
                        outtmpl = os.path.join(temp_dir, f'yt_{video_id}.%(ext)s')
                        
                        ydl_opts = {
                            'format': 'bestaudio/best',
                            'outtmpl': outtmpl,
                            'quiet': True,
                            'no_warnings': True,
                        }
                        
                        with yt_dlp.YoutubeDL(ydl_opts) as ydl:
                            info = ydl.extract_info(url, download=True)
                            filename = ydl.prepare_filename(info)
                        
                        if not os.path.exists(filename):
                            matched_files = glob.glob(os.path.join(temp_dir, f'yt_{video_id}.*'))
                            if matched_files:
                                filename = matched_files[0]
                            else:
                                raise Exception("Downloaded audio file not found on disk.")

                        try:
                            media_file = genai_client.files.upload(file=filename)
                            
                            while media_file.state.name == "PROCESSING":
                                time.sleep(1)
                                media_file = genai_client.files.get(name=media_file.name)
                            
                            if media_file.state.name == "FAILED":
                                raise Exception("Audio file processing failed on Gemini.")

                            response = genai_client.models.generate_content(
                                model='gemini-2.5-flash',
                                contents=[media_file, "Please transcribe the audio/dialogue in this video as accurately and completely as possible."]
                            )
                            transcript = response.text or ""
                            
                            genai_client.files.delete(name=media_file.name)
                            return jsonify({'success': True, 'transcript': transcript, 'calls_made': getattr(g, 'gemini_calls_count', 0)})
                        finally:
                            if os.path.exists(filename):
                                os.remove(filename)
                    except Exception as dl_err:
                        print(f"yt-dlp fallback failed: {dl_err}")
                        return jsonify({'error': (
                            f"Could not import YouTube video transcript. Direct Gemini processing error: {gemini_err}. "
                            f"Downloading error: {dl_err}"
                        )}), 400
        
        else:
            headers = {"User-Agent": "Mozilla/5.0"}
            r = requests.get(url, headers=headers, stream=True, timeout=30)
            r.raise_for_status()

            with tempfile.NamedTemporaryFile(delete=False, suffix='.mp4') as temp_file:
                for chunk in r.iter_content(chunk_size=8192):
                    if chunk:
                        temp_file.write(chunk)
                temp_path = temp_file.name

            try:
                media_file = genai_client.files.upload(file=temp_path)
                
                while media_file.state.name == "PROCESSING":
                    time.sleep(1)
                    media_file = genai_client.files.get(name=media_file.name)
                
                if media_file.state.name == "FAILED":
                    raise Exception("File processing failed on Gemini.")

                response = genai_client.models.generate_content(
                    model='gemini-2.5-flash',
                    contents=[media_file, "Please transcribe the audio/dialogue in this video as accurately and completely as possible."]
                )
                transcript = response.text or ""
                
                genai_client.files.delete(name=media_file.name)
                return jsonify({'success': True, 'transcript': transcript, 'calls_made': getattr(g, 'gemini_calls_count', 0)})
            finally:
                if os.path.exists(temp_path):
                    os.remove(temp_path)

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/pomodoro/log', methods=['POST'])
def log_pomodoro():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
        
    data = request.get_json() or {}
    category = data.get('category', 'School').strip()
    duration_minutes = data.get('duration_minutes', 25)
    
    if not category:
        return jsonify({'error': 'Category is required'}), 400
        
    db = get_db()
    try:
        db.execute('''
            INSERT INTO pomodoro_logs (user_id, category, duration_minutes)
            VALUES (?, ?, ?)
        ''', (session['user_id'], category, duration_minutes))
        db.commit()
        return jsonify({'success': True}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/pomodoro/stats', methods=['GET'])
def get_pomodoro_stats():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
        
    import datetime
    month_str = request.args.get('month', datetime.datetime.now().strftime('%Y-%m'))
    
    db = get_db()
    try:
        # Fetch pomodoro logs for the month
        p_logs = db.execute('''
            SELECT category, duration_minutes, completed_at
            FROM pomodoro_logs
            WHERE user_id = ? AND completed_at LIKE ?
            ORDER BY completed_at DESC
        ''', (session['user_id'], f"{month_str}%")).fetchall()
        
        pomodoro_list = []
        for r in p_logs:
            pomodoro_list.append({
                'category': r['category'],
                'duration_minutes': r['duration_minutes'],
                'completed_at': r['completed_at']
            })
            
        # Fetch completed tasks for the month
        t_completed = db.execute('''
            SELECT id, title, category, completed_at
            FROM tasks
            WHERE user_id = ? AND completed = 1 AND completed_at LIKE ?
            ORDER BY completed_at DESC
        ''', (session['user_id'], f"{month_str}%")).fetchall()
        
        tasks_list = []
        for r in t_completed:
            tasks_list.append({
                'id': r['id'],
                'title': r['title'],
                'category': r['category'],
                'completed_at': r['completed_at']
            })
            
        return jsonify({
            'pomodoro_logs': pomodoro_list,
            'completed_tasks': tasks_list
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/pomodoro/compile-report', methods=['GET', 'POST'])
def compile_report():
    import datetime
    from fpdf import FPDF
    import tempfile
    
    # 1. Determine user type and get data
    if request.method == 'POST':
        # Guest mode (data sent in body)
        data = request.get_json() or {}
        month_str = data.get('month', datetime.datetime.now().strftime('%Y-%m'))
        p_logs = data.get('pomodoro_logs', [])
        t_completed = data.get('completed_tasks', [])
    else:
        # Logged-in mode
        if 'user_id' not in session:
            return jsonify({'error': 'Unauthorized'}), 401
        
        month_str = request.args.get('month', datetime.datetime.now().strftime('%Y-%m'))
        
        db = get_db()
        p_rows = db.execute('''
            SELECT category, duration_minutes, completed_at
            FROM pomodoro_logs
            WHERE user_id = ? AND completed_at LIKE ?
            ORDER BY completed_at ASC
        ''', (session['user_id'], f"{month_str}%")).fetchall()
        
        p_logs = []
        for r in p_rows:
            p_logs.append({
                'category': r['category'],
                'duration_minutes': r['duration_minutes'],
                'completed_at': r['completed_at']
            })
            
        t_rows = db.execute('''
            SELECT title, category, completed_at
            FROM tasks
            WHERE user_id = ? AND completed = 1 AND completed_at LIKE ?
            ORDER BY completed_at ASC
        ''', (session['user_id'], f"{month_str}%")).fetchall()
        
        t_completed = []
        for r in t_rows:
            t_completed.append({
                'title': r['title'],
                'category': r['category'],
                'completed_at': r['completed_at']
            })
            
    # Parse month name
    try:
        dt = datetime.datetime.strptime(month_str, '%Y-%m')
        month_name = dt.strftime('%B %Y')
    except:
        month_name = month_str
        
    # 2. Compute stats
    category_minutes = {}
    total_minutes = 0
    for r in p_logs:
        cat = r.get('category', 'School')
        mins = int(r.get('duration_minutes', 25))
        category_minutes[cat] = category_minutes.get(cat, 0) + mins
        total_minutes += mins
        
    total_hours = round(total_minutes / 60.0, 1)
    
    # 3. Create PDF
    pdf = FPDF()
    pdf.alias_nb_pages()
    pdf.add_page()
    
    # Title
    pdf.set_font('helvetica', 'B', 20)
    pdf.set_text_color(46, 125, 50)
    pdf.cell(0, 15, "Lily's Study Pond Report", ln=True, align='C')
    pdf.set_font('helvetica', 'B', 12)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(0, 10, f"Activity & Productivity Report: {month_name}", ln=True, align='C')
    pdf.ln(10)
    
    # Overview
    pdf.set_font('helvetica', 'B', 14)
    pdf.set_text_color(33, 33, 33)
    pdf.cell(0, 8, "Overview", ln=True)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(3)
    
    pdf.set_font('helvetica', '', 11)
    pdf.cell(100, 8, f"Total Focus Hours: {total_hours} hrs", ln=False)
    pdf.cell(0, 8, f"Total Pomodoro Sessions: {len(p_logs)}", ln=True)
    pdf.cell(100, 8, f"Tasks Completed: {len(t_completed)}", ln=False)
    pdf.cell(0, 8, f"Report Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=True)
    pdf.ln(8)
    
    # Category Breakdown
    pdf.set_font('helvetica', 'B', 14)
    pdf.cell(0, 8, "Focus Time by Category", ln=True)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(3)
    
    pdf.set_font('helvetica', 'B', 11)
    pdf.cell(80, 8, "Category", border=1)
    pdf.cell(50, 8, "Focus Time (minutes)", border=1)
    pdf.cell(50, 8, "Focus Time (hours)", border=1, ln=True)
    
    pdf.set_font('helvetica', '', 11)
    if not category_minutes:
        pdf.cell(180, 8, "No focus sessions logged in this category/month.", border=1, ln=True, align='C')
    else:
        for cat, mins in sorted(category_minutes.items(), key=lambda x: x[1], reverse=True):
            hrs = round(mins / 60.0, 1)
            pdf.cell(80, 8, cat, border=1)
            pdf.cell(50, 8, str(mins), border=1)
            pdf.cell(50, 8, f"{hrs} hrs", border=1, ln=True)
    pdf.ln(10)
    
    # Tasks Completed
    pdf.set_font('helvetica', 'B', 14)
    pdf.cell(0, 8, f"Tasks Completed in {month_name}", ln=True)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(3)
    
    pdf.set_font('helvetica', 'B', 11)
    pdf.cell(100, 8, "Task Description", border=1)
    pdf.cell(40, 8, "Category", border=1)
    pdf.cell(40, 8, "Completed At", border=1, ln=True)
    
    pdf.set_font('helvetica', '', 10)
    if not t_completed:
        pdf.cell(180, 8, "No tasks completed in this month.", border=1, ln=True, align='C')
    else:
        for r in t_completed:
            title = r.get('title', '')
            cat = r.get('category', 'School')
            comp_time = r.get('completed_at', '') or "N/A"
            if len(title) > 45:
                title = title[:42] + "..."
            pdf.cell(100, 8, title, border=1)
            pdf.cell(40, 8, cat, border=1)
            pdf.cell(40, 8, comp_time[:16], border=1, ln=True)
    pdf.ln(10)
    
    # Detailed Activity Log
    pdf.set_font('helvetica', 'B', 14)
    pdf.cell(0, 8, "Focus Session Log", ln=True)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(3)
    
    pdf.set_font('helvetica', 'B', 11)
    pdf.cell(70, 8, "Completed Time", border=1)
    pdf.cell(70, 8, "Category", border=1)
    pdf.cell(40, 8, "Duration", border=1, ln=True)
    
    pdf.set_font('helvetica', '', 10)
    if not p_logs:
        pdf.cell(180, 8, "No focus sessions recorded.", border=1, ln=True, align='C')
    else:
        for r in p_logs:
            comp_time = r.get('completed_at', '')
            cat = r.get('category', 'School')
            duration = f"{r.get('duration_minutes', 25)} min"
            pdf.cell(70, 8, comp_time[:19], border=1)
            pdf.cell(70, 8, cat, border=1)
            pdf.cell(40, 8, duration, border=1, ln=True)
            
    temp_pdf = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
    pdf.output(temp_pdf.name)
    temp_pdf.close()
    
    filename = f"FrogGPT_Report_{month_str}.pdf"
    return send_file(temp_pdf.name, as_attachment=True, download_name=filename)

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5001)
